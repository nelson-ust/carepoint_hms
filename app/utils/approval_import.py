# app/utils/approval_import.py
"""
Excel helpers for bulk-importing the approval-engine configuration:
Request Types, Approval Flows and Approval Steps — each as its own
downloadable template + parser (uploaded in dependency order).

Mirrors ``app/utils/drug_import.py`` conventions:
- ``build_*_template`` — styled .xlsx (header row, dropdowns, Reference +
  Instructions sheets).
- ``parse_*_rows`` — read an uploaded workbook back into
  ``[{"row": <1-based row>, "data": {canonical_key: value}}]`` dicts,
  tolerant of column reordering.

Per product decision, Steps support only ROLE and DYNAMIC approvers.
"""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

YES_NO = ["YES", "NO"]
APPROVER_KINDS = ["ROLE", "DYNAMIC"]
DYNAMIC_TOKENS = [
    "REQUESTER_MANAGER", "DEPARTMENT_HEAD", "FACILITY_HEAD", "HR_HEAD", "FINANCE_HEAD",
]
DECISION_RULES = ["ANY_OF", "ALL_OF", "N_OF_M"]

_HEADER_FILL = PatternFill("solid", fgColor="10B981")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_TITLE_FONT = Font(bold=True, size=14, color="0F172A")
_HEAD_FONT = Font(bold=True, size=11, color="10B981")
_MAX_DATA_ROWS = 1000


def _clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value


def _write_header(ws, columns: list[tuple[str, str, bool, str]]) -> None:
    for col_idx, (header, _key, required, _help) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{header}{' *' if required else ''}")
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = max(16, len(header) + 6)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _col_letter(columns: list[tuple[str, str, bool, str]], key: str) -> str:
    for idx, (_h, k, _r, _help) in enumerate(columns, start=1):
        if k == key:
            return get_column_letter(idx)
    raise KeyError(key)


def _inline_dropdown(ws, columns, key, values, *, title, prompt) -> None:
    letter = _col_letter(columns, key)
    dv = DataValidation(
        type="list", formula1='"' + ",".join(values) + '"',
        allow_blank=True, showErrorMessage=True,
    )
    dv.error = "Pick a value from the list: " + ", ".join(values)
    dv.errorTitle = title
    dv.prompt = prompt
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{_MAX_DATA_ROWS + 1}")


def _ref_dropdown(ws, columns, key, *, ref_range, title, prompt) -> None:
    letter = _col_letter(columns, key)
    dv = DataValidation(type="list", formula1=ref_range, allow_blank=True, showErrorMessage=True)
    dv.error = "Pick a value from the Reference sheet."
    dv.errorTitle = title
    dv.prompt = prompt
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{_MAX_DATA_ROWS + 1}")


def _build_instructions(ws, title: str, steps: list[str], columns: list[tuple[str, str, bool, str]]) -> None:
    ws["A1"] = title
    ws["A1"].font = _TITLE_FONT
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 78
    row = 3
    ws.cell(row=row, column=1, value="Steps").font = _HEAD_FONT
    for line in steps:
        row += 1
        ws.cell(row=row, column=2, value=line)
    row += 2
    ws.cell(row=row, column=1, value="Columns").font = _HEAD_FONT
    for header, _key, required, help_text in columns:
        row += 1
        label = ws.cell(row=row, column=1, value=f"{header}{' *' if required else ''}")
        label.font = Font(bold=required)
        ws.cell(row=row, column=2, value=help_text)


def _parse(file_bytes: bytes, sheet_name: str, aliases: dict[str, str], required_key: str, what: str) -> list[dict[str, Any]]:
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=False)
    try:
        header_cells = next(rows_iter)
    except StopIteration:
        return []

    col_map: dict[int, str] = {}
    for cell in header_cells:
        if cell.value is None:
            continue
        text = str(cell.value).strip().lower().rstrip("*").strip()
        key = aliases.get(text)
        if key:
            col_map[cell.column] = key
    if required_key not in col_map.values():
        raise ValueError(f"The sheet doesn't look like the {what} template (missing a required column).")

    parsed: list[dict[str, Any]] = []
    for cells in rows_iter:
        data: dict[str, Any] = {}
        for cell in cells:
            key = col_map.get(cell.column)
            if not key:
                continue
            cleaned = _clean_cell(cell.value)
            if cleaned is not None:
                data[key] = cleaned
        if not data:
            continue
        parsed.append({"row": cells[0].row if cells else None, "data": data})
    return parsed


# ======================================================================
# 1) Request Types
# ======================================================================

REQUEST_TYPE_COLUMNS: list[tuple[str, str, bool, str]] = [
    ("Code", "code", True, "Unique code, UPPER_SNAKE_CASE (e.g. PROCUREMENT). Letters, digits, underscores."),
    ("Name", "name", True, "Human-friendly name (e.g. Procurement Request)."),
    ("Description", "description", False, "Optional description."),
    ("Active", "is_active", False, "YES (default) or NO."),
]

_REQUEST_TYPE_ALIASES = {
    "code": "code", "request type code": "code", "type code": "code",
    "name": "name", "request type name": "name",
    "description": "description", "desc": "description",
    "active": "is_active", "is active": "is_active",
}


def build_request_type_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Request Types"
    _write_header(ws, REQUEST_TYPE_COLUMNS)
    ws.column_dimensions["C"].width = 50
    _inline_dropdown(ws, REQUEST_TYPE_COLUMNS, "is_active", YES_NO, title="Invalid value", prompt="Active? YES/NO")
    _build_instructions(
        wb.create_sheet("Instructions"),
        "How to bulk-upload request types",
        [
            "1. Fill one row per request type on the 'Request Types' sheet.",
            "2. Code (*) and Name (*) are required. Code must be unique.",
            "3. Codes are stored in UPPER_SNAKE_CASE.",
            "4. Save the file, then upload it back in the app.",
            "5. Any rows with problems are reported by row number — fix and re-upload only those.",
        ],
        REQUEST_TYPE_COLUMNS,
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_request_type_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    return _parse(file_bytes, "Request Types", _REQUEST_TYPE_ALIASES, "code", "request type")


# ======================================================================
# 2) Flows
# ======================================================================

FLOW_COLUMNS: list[tuple[str, str, bool, str]] = [
    ("Request Type Code", "request_type", True, "Must match an existing request type code (see dropdown / Reference)."),
    ("Flow Code", "code", True, "Unique code within the request type (e.g. STANDARD)."),
    ("Flow Name", "name", True, "Human-friendly name (e.g. Standard Procurement Flow)."),
    ("Description", "description", False, "Optional description."),
    ("Default", "is_default", False, "YES to make this the default flow for the request type; otherwise NO."),
    ("Active", "is_active", False, "YES (default) or NO."),
]

_FLOW_ALIASES = {
    "request type code": "request_type", "request type": "request_type", "type code": "request_type",
    "flow code": "code", "code": "code",
    "flow name": "name", "name": "name",
    "description": "description", "desc": "description",
    "default": "is_default", "is default": "is_default",
    "active": "is_active", "is active": "is_active",
}


def build_flow_template(request_types: list[tuple[str, str]]) -> bytes:
    """request_types: list of (code, name) for the tenant's existing request types."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Flows"
    _write_header(ws, FLOW_COLUMNS)
    ws.column_dimensions["D"].width = 46

    _inline_dropdown(ws, FLOW_COLUMNS, "is_default", YES_NO, title="Invalid value", prompt="Default flow? YES/NO")
    _inline_dropdown(ws, FLOW_COLUMNS, "is_active", YES_NO, title="Invalid value", prompt="Active? YES/NO")

    ref = wb.create_sheet("Reference")
    ref["A1"] = "Request Type Code"
    ref["B1"] = "Request Type Name"
    for c in ("A1", "B1"):
        ref[c].font = Font(bold=True)
    for i, (code, name) in enumerate(request_types, start=2):
        ref.cell(row=i, column=1, value=code)
        ref.cell(row=i, column=2, value=name)
    ref.column_dimensions["A"].width = 28
    ref.column_dimensions["B"].width = 30
    if request_types:
        last = len(request_types) + 1
        _ref_dropdown(ws, FLOW_COLUMNS, "request_type", ref_range=f"=Reference!$A$2:$A${last}",
                      title="Unknown Request Type", prompt="Choose an existing request type code")

    _build_instructions(
        wb.create_sheet("Instructions"),
        "How to bulk-upload approval flows",
        [
            "1. Create the Request Types first (separate template), then fill this sheet.",
            "2. Request Type Code (*), Flow Code (*) and Flow Name (*) are required.",
            "3. Request Type Code must match an existing request type (see Reference / dropdown).",
            "4. Flow Code must be unique within its request type.",
            "5. Set Default = YES on at most one flow per request type (it replaces any current default).",
            "6. Save the file, then upload it back in the app. Add the steps afterwards with the Steps template.",
        ],
        FLOW_COLUMNS,
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_flow_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    return _parse(file_bytes, "Flows", _FLOW_ALIASES, "code", "flow")


# ======================================================================
# 3) Steps
# ======================================================================

STEP_COLUMNS: list[tuple[str, str, bool, str]] = [
    ("Request Type Code", "request_type", True, "Request type code the flow belongs to (see Reference)."),
    ("Flow Code", "flow_code", True, "Flow code within that request type (see Reference)."),
    ("Step Order", "step_order", True, "Positive integer; unique within the flow (1, 2, 3 …)."),
    ("Step Name", "name", True, "Human-friendly step name (e.g. HOD Review)."),
    ("Approver Kind", "approver_kind", True, "ROLE or DYNAMIC."),
    ("Role", "approver_role", False, "For ROLE: an existing role code or name (see Reference)."),
    ("Dynamic Token", "dynamic_token", False, "For DYNAMIC: one of " + ", ".join(DYNAMIC_TOKENS) + "."),
    ("Decision Rule", "decision_rule", False, "ANY_OF (default), ALL_OF or N_OF_M."),
    ("Required Approvals", "required_approvals", False, "Number of approvals needed (used with N_OF_M). Default 1."),
    ("Allow Self Approval", "allow_self_approval", False, "YES or NO (default NO)."),
    ("Active", "is_active", False, "YES (default) or NO."),
]

_STEP_ALIASES = {
    "request type code": "request_type", "request type": "request_type", "type code": "request_type",
    "flow code": "flow_code", "flow": "flow_code",
    "step order": "step_order", "order": "step_order",
    "step name": "name", "name": "name",
    "approver kind": "approver_kind", "kind": "approver_kind",
    "role": "approver_role", "approver role": "approver_role",
    "dynamic token": "dynamic_token", "token": "dynamic_token",
    "decision rule": "decision_rule", "rule": "decision_rule",
    "required approvals": "required_approvals", "required": "required_approvals",
    "allow self approval": "allow_self_approval", "self approval": "allow_self_approval",
    "active": "is_active", "is active": "is_active",
}


def build_step_template(flows: list[tuple[str, str, str]], roles: list[tuple[str, str]]) -> bytes:
    """
    flows: list of (request_type_code, flow_code, flow_name).
    roles: list of (code, name).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Steps"
    _write_header(ws, STEP_COLUMNS)

    _inline_dropdown(ws, STEP_COLUMNS, "approver_kind", APPROVER_KINDS, title="Invalid Approver Kind", prompt="ROLE or DYNAMIC")
    _inline_dropdown(ws, STEP_COLUMNS, "dynamic_token", DYNAMIC_TOKENS, title="Invalid Token", prompt="Pick a dynamic approver token")
    _inline_dropdown(ws, STEP_COLUMNS, "decision_rule", DECISION_RULES, title="Invalid Rule", prompt="ANY_OF / ALL_OF / N_OF_M")
    _inline_dropdown(ws, STEP_COLUMNS, "allow_self_approval", YES_NO, title="Invalid value", prompt="Allow self approval? YES/NO")
    _inline_dropdown(ws, STEP_COLUMNS, "is_active", YES_NO, title="Invalid value", prompt="Active? YES/NO")

    ref = wb.create_sheet("Reference")
    ref["A1"] = "Request Type Code"
    ref["B1"] = "Flow Code"
    ref["C1"] = "Flow Name"
    ref["E1"] = "Role Code"
    ref["F1"] = "Role Name"
    for c in ("A1", "B1", "C1", "E1", "F1"):
        ref[c].font = Font(bold=True)
    for i, (rt_code, flow_code, flow_name) in enumerate(flows, start=2):
        ref.cell(row=i, column=1, value=rt_code)
        ref.cell(row=i, column=2, value=flow_code)
        ref.cell(row=i, column=3, value=flow_name)
    for i, (code, name) in enumerate(roles, start=2):
        ref.cell(row=i, column=5, value=code)
        ref.cell(row=i, column=6, value=name)
    for col, w in (("A", 26), ("B", 22), ("C", 30), ("E", 24), ("F", 28)):
        ref.column_dimensions[col].width = w

    if flows:
        last_f = len(flows) + 1
        _ref_dropdown(ws, STEP_COLUMNS, "request_type", ref_range=f"=Reference!$A$2:$A${last_f}",
                      title="Unknown Request Type", prompt="Request type of the flow")
        _ref_dropdown(ws, STEP_COLUMNS, "flow_code", ref_range=f"=Reference!$B$2:$B${last_f}",
                      title="Unknown Flow", prompt="Flow code (must match its request type)")
    if roles:
        last_r = len(roles) + 1
        _ref_dropdown(ws, STEP_COLUMNS, "approver_role", ref_range=f"=Reference!$E$2:$E${last_r}",
                      title="Unknown Role", prompt="Role code (used when Approver Kind = ROLE)")

    _build_instructions(
        wb.create_sheet("Instructions"),
        "How to bulk-upload approval steps",
        [
            "1. Create the Request Types and Flows first (separate templates), then fill this sheet.",
            "2. Request Type Code (*), Flow Code (*), Step Order (*), Step Name (*) and Approver Kind (*) are required.",
            "3. Approver Kind = ROLE → fill the Role column (existing role code/name).",
            "4. Approver Kind = DYNAMIC → fill the Dynamic Token column instead.",
            "5. Step Order must be a unique positive integer within its flow.",
            "6. Save the file, then upload it back in the app.",
            "7. Any rows with problems are reported by row number — fix and re-upload only those.",
        ],
        STEP_COLUMNS,
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_step_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    return _parse(file_bytes, "Steps", _STEP_ALIASES, "flow_code", "step")
