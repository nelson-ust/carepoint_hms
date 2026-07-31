# app/utils/account_import.py
"""
Excel helpers for bulk-importing the chart of accounts and billable services.

Mirrors ``app/utils/drug_import.py``:
- ``build_account_template`` / ``parse_account_rows``
- ``build_billable_service_template`` / ``parse_billable_service_rows``

Each ``parse_*`` returns ``[{"row": <1-based sheet row>, "data": {key: value}}]``
and is tolerant of column reordering via header aliases.
"""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

_HEADER_FILL = PatternFill("solid", fgColor="10B981")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_MAX_DATA_ROWS = 1000

ACCOUNT_TYPES = ["REVENUE", "ASSET", "LIABILITY", "EXPENSE", "EQUITY"]


def _clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value


def _write_header(ws, columns: list[tuple[str, str, bool, str]]) -> None:
    for col_idx, (header, _key, required, _help) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{header}{' *' if required else ''}")
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = max(18, len(header) + 8)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _build_instructions(ws, title: str, steps: list[str], columns: list[tuple[str, str, bool, str]]) -> None:
    title_font = Font(bold=True, size=14, color="0F172A")
    head_font = Font(bold=True, size=11, color="10B981")
    ws["A1"] = title
    ws["A1"].font = title_font
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 74
    row = 3
    ws.cell(row=row, column=1, value="Steps").font = head_font
    for line in steps:
        row += 1
        ws.cell(row=row, column=2, value=line)
    row += 2
    ws.cell(row=row, column=1, value="Columns").font = head_font
    for header, _key, required, help_text in columns:
        row += 1
        label = ws.cell(row=row, column=1, value=f"{header}{' *' if required else ''}")
        label.font = Font(bold=required)
        ws.cell(row=row, column=2, value=help_text)


def _parse_rows(file_bytes: bytes, sheet_name: str, aliases: dict[str, str], required_key: str, what: str) -> list[dict[str, Any]]:
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=False)
    try:
        header_cells = next(rows_iter)
    except StopIteration:
        return []

    col_map: dict[int, str] = {}
    for cell in header_cells:
        if cell.value is None:
            continue
        text = str(cell.value).strip().lower().rstrip("*").strip()
        key = aliases.get(text)
        if key:
            col_map[cell.column] = key
    if required_key not in col_map.values():
        raise ValueError(f"The sheet doesn't look like the {what} template.")

    parsed: list[dict[str, Any]] = []
    for cells in rows_iter:
        data: dict[str, Any] = {}
        for cell in cells:
            key = col_map.get(cell.column)
            if not key:
                continue
            cleaned = _clean_cell(cell.value)
            if cleaned is not None:
                data[key] = cleaned
        if not data:
            continue
        parsed.append({"row": cells[0].row if cells else None, "data": data})
    return parsed


# ======================================================================
# Chart of accounts
# ======================================================================

ACCOUNT_COLUMNS: list[tuple[str, str, bool, str]] = [
    ("Account Code", "code", True, "Short unique code (e.g. REV-CONS)."),
    ("Account Name", "name", True, "Account name (e.g. Consultation Revenue)."),
    ("Account Type", "account_type", False, "One of: " + ", ".join(ACCOUNT_TYPES) + ". Defaults to REVENUE."),
    ("Description", "description", False, "Optional description."),
]

_ACCOUNT_HEADER_ALIASES: dict[str, str] = {
    "account code": "code", "code": "code",
    "account name": "name", "name": "name",
    "account type": "account_type", "type": "account_type",
    "description": "description", "desc": "description",
}


def build_account_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Accounts"
    _write_header(ws, ACCOUNT_COLUMNS)
    ws.column_dimensions["D"].width = 50

    type_letter = get_column_letter(3)  # Account Type column
    dv = DataValidation(type="list", formula1='"' + ",".join(ACCOUNT_TYPES) + '"', allow_blank=True, showErrorMessage=True)
    dv.error = "Pick a value: " + ", ".join(ACCOUNT_TYPES)
    dv.errorTitle = "Invalid Account Type"
    dv.prompt = "Choose the account type"
    ws.add_data_validation(dv)
    dv.add(f"{type_letter}2:{type_letter}{_MAX_DATA_ROWS + 1}")

    _build_instructions(
        wb.create_sheet("Instructions"),
        "How to bulk-upload the chart of accounts",
        [
            "1. Fill one row per account on the 'Accounts' sheet.",
            "2. Account Code (*) and Account Name (*) are required; Code must be unique.",
            "3. Account Type has a dropdown — defaults to REVENUE when blank.",
            "4. Save the file, then upload it back in the app.",
            "5. Rows with problems are reported by row number — fix and re-upload only those.",
        ],
        ACCOUNT_COLUMNS,
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_account_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    return _parse_rows(file_bytes, "Accounts", _ACCOUNT_HEADER_ALIASES, "code", "account")


# ======================================================================
# Billable services
# ======================================================================

SERVICE_COLUMNS: list[tuple[str, str, bool, str]] = [
    ("Service Code", "code", True, "Short unique code (e.g. CONSULT_GP)."),
    ("Service Name", "name", True, "Service name."),
    ("Category", "category", False, "Optional grouping (e.g. CONSULTATION)."),
    ("Rate", "default_price", True, "Amount charged; must be greater than 0."),
    ("Account Code", "account_code", True, "Must match an existing account code (see Reference)."),
    ("Description", "description", False, "Optional description."),
]

_SERVICE_HEADER_ALIASES: dict[str, str] = {
    "service code": "code", "code": "code",
    "service name": "name", "name": "name",
    "category": "category",
    "rate": "default_price", "price": "default_price", "default price": "default_price", "amount": "default_price",
    "account code": "account_code", "account": "account_code",
    "description": "description", "desc": "description",
}


def build_billable_service_template(accounts: list[tuple[str, str]]) -> bytes:
    """``accounts`` is a list of (code, name) for the tenant's existing accounts."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Services"
    _write_header(ws, SERVICE_COLUMNS)
    ws.column_dimensions["F"].width = 50

    ref = wb.create_sheet("Reference")
    ref["A1"] = "Account Code"
    ref["B1"] = "Account Name"
    for c in ("A1", "B1"):
        ref[c].font = Font(bold=True)
    for i, (code, name) in enumerate(accounts, start=2):
        ref.cell(row=i, column=1, value=code)
        ref.cell(row=i, column=2, value=name)
    ref.column_dimensions["A"].width = 22
    ref.column_dimensions["B"].width = 32

    if accounts:
        acc_letter = get_column_letter(5)  # Account Code column
        last_ref = len(accounts) + 1
        dv = DataValidation(type="list", formula1=f"Reference!$A$2:$A${last_ref}", allow_blank=True, showErrorMessage=True)
        dv.error = "Pick an existing account code (see the Reference sheet)."
        dv.errorTitle = "Unknown Account"
        dv.prompt = "Choose the posting account"
        ws.add_data_validation(dv)
        dv.add(f"{acc_letter}2:{acc_letter}{_MAX_DATA_ROWS + 1}")

    ins = wb.create_sheet("Instructions")
    _build_instructions(
        ins,
        "How to bulk-upload billable services",
        [
            "1. Fill one row per service on the 'Services' sheet.",
            "2. Service Code (*), Service Name (*), Rate (*) and Account Code (*) are required.",
            "3. Rate must be greater than 0.",
            "4. Account Code must match one in the Reference sheet (dropdown provided).",
            "5. Save the file, then upload it back in the app.",
            "6. Rows with problems are reported by row number — fix and re-upload only those.",
        ],
        SERVICE_COLUMNS,
    )
    row = 3 + len(SERVICE_COLUMNS) + 4
    ins.cell(row=row, column=1, value="Valid accounts").font = Font(bold=True, size=11, color="10B981")
    if accounts:
        for code, name in accounts:
            row += 1
            ins.cell(row=row, column=1, value=code).font = Font(bold=True)
            ins.cell(row=row, column=2, value=name)
    else:
        row += 1
        ins.cell(row=row, column=2, value="No accounts yet — create the chart of accounts first, then re-download this template.")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_billable_service_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    return _parse_rows(file_bytes, "Services", _SERVICE_HEADER_ALIASES, "code", "billable service")
