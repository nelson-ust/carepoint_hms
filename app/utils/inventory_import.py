# app/utils/inventory_import.py
"""
Excel helpers for bulk stock-item import.

Provides:
- ``build_stock_item_template`` — generate a styled .xlsx with a pre-populated
  Item Type dropdown (and a Store Code dropdown sourced from the tenant's own
  stores) plus an Instructions sheet.
- ``parse_stock_item_rows`` — read an uploaded workbook back into a list of
  row dicts keyed by our canonical field names, tolerant of column reordering.
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Authoritative item types — must match the InventoryItemType enum /
# InventoryStockItemCreateSchema validator on the backend.
ITEM_TYPES = ["DRUG", "CONSUMABLE", "EQUIPMENT", "SUPPLY", "OTHER"]

# Ordered template columns: (header, canonical_key, required, help text)
TEMPLATE_COLUMNS: list[tuple[str, str, bool, str]] = [
    ("Store Code", "store_code", True, "Must match an existing store code (see dropdown / Instructions)."),
    ("Item Type", "item_type", True, "One of: " + ", ".join(ITEM_TYPES) + "."),
    ("Drug (SKU or Name)", "drug_ref", False, "Medicines only: link to an existing drug by SKU or exact name. Auto-fills the item name / unit for DRUG rows."),
    ("Item Name", "item_name", True, "Descriptive name. For a linked drug you may leave this blank — it is taken from the drug."),
    ("SKU", "sku", False, "Stock keeping unit (optional)."),
    ("Unit of Measure", "unit_of_measure", False, "e.g. Tablet, Vial, Box, Pack."),
    ("Quantity on Hand", "quantity_on_hand", True, "Opening balance — a number (e.g. 100)."),
    ("Reorder Level", "reorder_level", False, "Low-stock alert threshold (number)."),
    ("Unit Cost", "unit_cost", False, "Procurement cost per unit (number)."),
    ("Batch No", "batch_no", False, "Manufacturer batch number (optional)."),
    ("Expiry Date", "expiry_date", False, "Format YYYY-MM-DD (e.g. 2027-01-31)."),
]

# Header (lower-cased, without trailing '*') -> canonical key. Aliases included
# so a slightly different header still maps correctly.
_HEADER_ALIASES: dict[str, str] = {
    "store code": "store_code",
    "store": "store_code",
    "item type": "item_type",
    "type": "item_type",
    "drug (sku or name)": "drug_ref",
    "drug": "drug_ref",
    "drug ref": "drug_ref",
    "drug sku": "drug_ref",
    "drug name": "drug_ref",
    "linked drug": "drug_ref",
    "item name": "item_name",
    "name": "item_name",
    "sku": "sku",
    "unit of measure": "unit_of_measure",
    "uom": "unit_of_measure",
    "quantity on hand": "quantity_on_hand",
    "quantity": "quantity_on_hand",
    "qty": "quantity_on_hand",
    "reorder level": "reorder_level",
    "reorder": "reorder_level",
    "unit cost": "unit_cost",
    "cost": "unit_cost",
    "batch no": "batch_no",
    "batch number": "batch_no",
    "batch": "batch_no",
    "expiry date": "expiry_date",
    "expiry": "expiry_date",
}

_HEADER_FILL = PatternFill("solid", fgColor="10B981")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_REQUIRED_MARK_FONT = Font(bold=True, color="FFFFFF", size=11)
_MAX_DATA_ROWS = 1000


def build_stock_item_template(stores: list[tuple[str, str]], drugs: list[tuple[str, str]] | None = None) -> bytes:
    """
    Build the bulk-upload template.

    Args:
        stores: list of (code, name) for the tenant's existing stores.

    Returns:
        The .xlsx file as bytes.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock Items"

    # Header row
    for col_idx, (header, _key, required, _help) in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{header}{' *' if required else ''}")
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = max(16, len(header) + 6)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    # Item Type dropdown (inline list — short enough for Excel's 255-char limit)
    type_col = _column_for_key("item_type")
    type_letter = get_column_letter(type_col)
    type_dv = DataValidation(
        type="list",
        formula1='"' + ",".join(ITEM_TYPES) + '"',
        allow_blank=False,
        showErrorMessage=True,
    )
    type_dv.error = "Pick a value from the list: " + ", ".join(ITEM_TYPES)
    type_dv.errorTitle = "Invalid Item Type"
    type_dv.prompt = "Choose the item type"
    ws.add_data_validation(type_dv)
    type_dv.add(f"{type_letter}2:{type_letter}{_MAX_DATA_ROWS + 1}")

    # Reference sheet holding store codes (drives the Store Code dropdown and
    # gives users a legend of valid codes).
    ref = wb.create_sheet("Reference")
    ref["A1"] = "Store Code"
    ref["B1"] = "Store Name"
    for c in ("A1", "B1"):
        ref[c].font = Font(bold=True)
    for i, (code, name) in enumerate(stores, start=2):
        ref.cell(row=i, column=1, value=code)
        ref.cell(row=i, column=2, value=name)
    ref.column_dimensions["A"].width = 20
    ref.column_dimensions["B"].width = 32

    if stores:
        store_col = _column_for_key("store_code")
        store_letter = get_column_letter(store_col)
        last_ref = len(stores) + 1
        store_dv = DataValidation(
            type="list",
            formula1=f"Reference!$A$2:$A${last_ref}",
            allow_blank=False,
            showErrorMessage=True,
        )
        store_dv.error = "Pick an existing store code (see the Reference sheet)."
        store_dv.errorTitle = "Unknown Store Code"
        store_dv.prompt = "Choose the destination store"
        ws.add_data_validation(store_dv)
        store_dv.add(f"{store_letter}2:{store_letter}{_MAX_DATA_ROWS + 1}")

    # Drug reference sheet + dropdown (drives the optional Drug link so
    # medicines are tied to the formulary instead of being re-typed).
    drug_rows = drugs or []
    if drug_rows:
        dref = wb.create_sheet("Drugs")
        dref["A1"] = "Drug Name"
        dref["B1"] = "SKU"
        for c in ("A1", "B1"):
            dref[c].font = Font(bold=True)
        for i, (dname, dsku) in enumerate(drug_rows, start=2):
            dref.cell(row=i, column=1, value=dname)
            dref.cell(row=i, column=2, value=dsku)
        dref.column_dimensions["A"].width = 36
        dref.column_dimensions["B"].width = 20

        drug_col = _column_for_key("drug_ref")
        drug_letter = get_column_letter(drug_col)
        last_drug = len(drug_rows) + 1
        # allow_blank + no hard error: users may also type a SKU not in the
        # name list; the importer resolves by name OR sku.
        drug_dv = DataValidation(
            type="list",
            formula1=f"Drugs!$A$2:$A${last_drug}",
            allow_blank=True,
            showErrorMessage=False,
        )
        drug_dv.prompt = "Optional: link a medicine to its formulary drug (by name or SKU)."
        ws.add_data_validation(drug_dv)
        drug_dv.add(f"{drug_letter}2:{drug_letter}{_MAX_DATA_ROWS + 1}")

    # Instructions sheet
    _build_instructions_sheet(wb.create_sheet("Instructions"), stores)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_instructions_sheet(ws, stores: list[tuple[str, str]]) -> None:
    title_font = Font(bold=True, size=14, color="0F172A")
    head_font = Font(bold=True, size=11, color="10B981")
    ws["A1"] = "How to bulk-upload stock items"
    ws["A1"].font = title_font
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 70

    row = 3
    ws.cell(row=row, column=1, value="Steps").font = head_font
    for line in (
        "1. Fill one row per stock item on the 'Stock Items' sheet.",
        "2. Store Code and Item Type have dropdowns — pick from the list.",
        "3. Columns marked with * are required.",
        "4. Save the file, then upload it back in the app.",
        "5. Any rows with problems are reported by row number — fix and re-upload only those.",
    ):
        row += 1
        ws.cell(row=row, column=2, value=line)

    row += 2
    ws.cell(row=row, column=1, value="Columns").font = head_font
    for header, _key, required, help_text in TEMPLATE_COLUMNS:
        row += 1
        label = ws.cell(row=row, column=1, value=f"{header}{' *' if required else ''}")
        label.font = Font(bold=required)
        ws.cell(row=row, column=2, value=help_text)

    row += 2
    ws.cell(row=row, column=1, value="Valid store codes").font = head_font
    if stores:
        for code, name in stores:
            row += 1
            ws.cell(row=row, column=1, value=code).font = Font(bold=True)
            ws.cell(row=row, column=2, value=name)
    else:
        row += 1
        ws.cell(row=row, column=2, value="No stores exist yet — create a store first, then download a fresh template.")


def _column_for_key(key: str) -> int:
    for idx, (_h, k, _r, _help) in enumerate(TEMPLATE_COLUMNS, start=1):
        if k == key:
            return idx
    raise KeyError(key)


def _normalize_header(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip().lower().rstrip("*").strip()
    return _HEADER_ALIASES.get(text)


def _clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    if isinstance(value, datetime):
        return value.date()
    return value


def parse_stock_item_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    """
    Parse an uploaded template into row dicts.

    Returns a list of ``{"row": <1-based sheet row>, "data": {canonical_key: value}}``,
    skipping fully-empty rows. Raises ValueError if the header row can't be
    understood.
    """
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb["Stock Items"] if "Stock Items" in wb.sheetnames else wb[wb.sheetnames[0]]

    rows_iter = ws.iter_rows(values_only=False)
    try:
        header_cells = next(rows_iter)
    except StopIteration:
        return []

    col_map: dict[int, str] = {}
    for cell in header_cells:
        key = _normalize_header(cell.value)
        if key:
            col_map[cell.column] = key
    if "store_code" not in col_map.values() or "item_name" not in col_map.values():
        raise ValueError(
            "The sheet doesn't look like the stock-item template "
            "(missing the 'Store Code' or 'Item Name' column)."
        )

    parsed: list[dict[str, Any]] = []
    for cells in rows_iter:
        data: dict[str, Any] = {}
        for cell in cells:
            key = col_map.get(cell.column)
            if not key:
                continue
            cleaned = _clean_cell(cell.value)
            if cleaned is not None:
                data[key] = cleaned
        if not data:
            continue  # skip blank rows
        row_number = cells[0].row if cells else None
        parsed.append({"row": row_number, "data": data})
    return parsed
