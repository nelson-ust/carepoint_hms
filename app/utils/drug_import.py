# app/utils/drug_import.py
"""
Excel helpers for bulk drug (formulary) import.

Mirrors ``app/utils/inventory_import.py``:
- ``build_drug_template`` — generate a styled .xlsx with pre-populated Dosage
  Form and Controlled dropdowns (and a Category dropdown sourced from the
  tenant's own drug categories) plus an Instructions sheet.
- ``parse_drug_rows`` — read an uploaded workbook back into a list of row dicts
  keyed by our canonical field names, tolerant of column reordering.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Common dosage forms — kept in sync with the frontend DOSAGE_FORMS list.
DOSAGE_FORMS = [
    "Tablet", "Capsule", "Syrup", "Suspension", "Injection", "Cream",
    "Ointment", "Drops", "Inhaler", "Patch", "Powder", "Other",
]

YES_NO = ["YES", "NO"]

# Ordered template columns: (header, canonical_key, required, help text)
TEMPLATE_COLUMNS: list[tuple[str, str, bool, str]] = [
    ("Drug Name", "name", True, "Full product / formulary name (e.g. Paracetamol 500mg Tablet)."),
    ("Generic Name", "generic_name", False, "International non-proprietary name (e.g. Paracetamol)."),
    ("Brand Name", "brand_name", False, "Manufacturer brand, if any (optional)."),
    ("Strength", "strength", False, "e.g. 500mg, 250mg/5ml."),
    ("Dosage Form", "dosage_form", False, "One of: " + ", ".join(DOSAGE_FORMS) + "."),
    ("Pack Size", "pack_size", False, "e.g. 10 x 10, 100ml bottle."),
    ("SKU", "sku", False, "Stock keeping unit — must be unique (optional)."),
    ("Category", "category", False, "Must match an existing category name/code (see dropdown)."),
    ("Unit Price", "unit_price", False, "Selling price per unit (number)."),
    ("Reorder Level", "reorder_level", False, "Low-stock alert threshold (number)."),
    ("Controlled", "is_controlled", False, "YES for a controlled substance, otherwise NO."),
]

# Header (lower-cased, without trailing '*') -> canonical key, with aliases.
_HEADER_ALIASES: dict[str, str] = {
    "drug name": "name",
    "name": "name",
    "generic name": "generic_name",
    "generic": "generic_name",
    "brand name": "brand_name",
    "brand": "brand_name",
    "strength": "strength",
    "dosage form": "dosage_form",
    "form": "dosage_form",
    "pack size": "pack_size",
    "pack": "pack_size",
    "sku": "sku",
    "category": "category",
    "drug category": "category",
    "unit price": "unit_price",
    "price": "unit_price",
    "reorder level": "reorder_level",
    "reorder": "reorder_level",
    "controlled": "is_controlled",
    "is controlled": "is_controlled",
    "controlled substance": "is_controlled",
}

_HEADER_FILL = PatternFill("solid", fgColor="10B981")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_MAX_DATA_ROWS = 1000


def _column_for_key(key: str) -> int:
    for idx, (_h, k, _r, _help) in enumerate(TEMPLATE_COLUMNS, start=1):
        if k == key:
            return idx
    raise KeyError(key)


def build_drug_template(categories: list[tuple[str, str | None]]) -> bytes:
    """
    Build the bulk drug-upload template.

    Args:
        categories: list of (name, code) for the tenant's existing drug
            categories (drives the Category dropdown / Reference legend).

    Returns:
        The .xlsx file as bytes.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Drugs"

    # Header row
    for col_idx, (header, _key, required, _help) in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{header}{' *' if required else ''}")
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = max(16, len(header) + 6)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    # Dosage Form dropdown (inline list)
    form_letter = get_column_letter(_column_for_key("dosage_form"))
    form_dv = DataValidation(
        type="list", formula1='"' + ",".join(DOSAGE_FORMS) + '"',
        allow_blank=True, showErrorMessage=True,
    )
    form_dv.error = "Pick a value from the list: " + ", ".join(DOSAGE_FORMS)
    form_dv.errorTitle = "Invalid Dosage Form"
    form_dv.prompt = "Choose the dosage form"
    ws.add_data_validation(form_dv)
    form_dv.add(f"{form_letter}2:{form_letter}{_MAX_DATA_ROWS + 1}")

    # Controlled dropdown (YES/NO)
    ctrl_letter = get_column_letter(_column_for_key("is_controlled"))
    ctrl_dv = DataValidation(
        type="list", formula1='"' + ",".join(YES_NO) + '"',
        allow_blank=True, showErrorMessage=True,
    )
    ctrl_dv.error = "Pick YES or NO."
    ctrl_dv.errorTitle = "Invalid value"
    ctrl_dv.prompt = "Is this a controlled substance?"
    ws.add_data_validation(ctrl_dv)
    ctrl_dv.add(f"{ctrl_letter}2:{ctrl_letter}{_MAX_DATA_ROWS + 1}")

    # Reference sheet holding category names (drives the Category dropdown).
    ref = wb.create_sheet("Reference")
    ref["A1"] = "Category Name"
    ref["B1"] = "Category Code"
    for c in ("A1", "B1"):
        ref[c].font = Font(bold=True)
    for i, (name, code) in enumerate(categories, start=2):
        ref.cell(row=i, column=1, value=name)
        ref.cell(row=i, column=2, value=code)
    ref.column_dimensions["A"].width = 28
    ref.column_dimensions["B"].width = 20

    if categories:
        cat_letter = get_column_letter(_column_for_key("category"))
        last_ref = len(categories) + 1
        cat_dv = DataValidation(
            type="list", formula1=f"Reference!$A$2:$A${last_ref}",
            allow_blank=True, showErrorMessage=True,
        )
        cat_dv.error = "Pick an existing category (see the Reference sheet)."
        cat_dv.errorTitle = "Unknown Category"
        cat_dv.prompt = "Choose the drug category (optional)"
        ws.add_data_validation(cat_dv)
        cat_dv.add(f"{cat_letter}2:{cat_letter}{_MAX_DATA_ROWS + 1}")

    _build_instructions_sheet(wb.create_sheet("Instructions"), categories)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_instructions_sheet(ws, categories: list[tuple[str, str | None]]) -> None:
    title_font = Font(bold=True, size=14, color="0F172A")
    head_font = Font(bold=True, size=11, color="10B981")
    ws["A1"] = "How to bulk-upload drugs"
    ws["A1"].font = title_font
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 72

    row = 3
    ws.cell(row=row, column=1, value="Steps").font = head_font
    for line in (
        "1. Fill one row per drug on the 'Drugs' sheet.",
        "2. Dosage Form, Category and Controlled have dropdowns — pick from the list.",
        "3. Columns marked with * are required.",
        "4. SKU (if provided) must be unique across the formulary.",
        "5. Save the file, then upload it back in the app.",
        "6. Any rows with problems are reported by row number — fix and re-upload only those.",
    ):
        row += 1
        ws.cell(row=row, column=2, value=line)

    row += 2
    ws.cell(row=row, column=1, value="Columns").font = head_font
    for header, _key, required, help_text in TEMPLATE_COLUMNS:
        row += 1
        label = ws.cell(row=row, column=1, value=f"{header}{' *' if required else ''}")
        label.font = Font(bold=required)
        ws.cell(row=row, column=2, value=help_text)

    row += 2
    ws.cell(row=row, column=1, value="Valid categories").font = head_font
    if categories:
        for name, code in categories:
            row += 1
            ws.cell(row=row, column=1, value=name).font = Font(bold=True)
            ws.cell(row=row, column=2, value=code or "")
    else:
        row += 1
        ws.cell(row=row, column=2, value="No categories yet — leave Category blank, or create categories first for a richer template.")


def _normalize_header(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip().lower().rstrip("*").strip()
    return _HEADER_ALIASES.get(text)


def _clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value


def parse_drug_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    """
    Parse an uploaded template into row dicts.

    Returns ``[{"row": <1-based sheet row>, "data": {canonical_key: value}}]``,
    skipping fully-empty rows. Raises ValueError if the header row can't be
    understood.
    """
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb["Drugs"] if "Drugs" in wb.sheetnames else wb[wb.sheetnames[0]]

    rows_iter = ws.iter_rows(values_only=False)
    try:
        header_cells = next(rows_iter)
    except StopIteration:
        return []

    col_map: dict[int, str] = {}
    for cell in header_cells:
        key = _normalize_header(cell.value)
        if key:
            col_map[cell.column] = key
    if "name" not in col_map.values():
        raise ValueError(
            "The sheet doesn't look like the drug template (missing the 'Drug Name' column)."
        )

    parsed: list[dict[str, Any]] = []
    for cells in rows_iter:
        data: dict[str, Any] = {}
        for cell in cells:
            key = col_map.get(cell.column)
            if not key:
                continue
            cleaned = _clean_cell(cell.value)
            if cleaned is not None:
                data[key] = cleaned
        if not data:
            continue
        row_number = cells[0].row if cells else None
        parsed.append({"row": row_number, "data": data})
    return parsed


# ======================================================================
# Drug-category bulk template
# ======================================================================

CATEGORY_COLUMNS: list[tuple[str, str, bool, str]] = [
    ("Category Name", "name", True, "Unique category name (e.g. Analgesics)."),
    ("Category Code", "code", False, "Short unique code (optional, e.g. ANALG)."),
    ("Description", "description", False, "Optional description of the category."),
]

_CATEGORY_HEADER_ALIASES: dict[str, str] = {
    "category name": "name",
    "name": "name",
    "category code": "code",
    "code": "code",
    "description": "description",
    "desc": "description",
}


def build_category_template() -> bytes:
    """Build the drug-category bulk-upload template (.xlsx bytes)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Categories"

    for col_idx, (header, _key, required, _help) in enumerate(CATEGORY_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=f"{header}{' *' if required else ''}")
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = max(20, len(header) + 8)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    ws.column_dimensions["C"].width = 50

    ins = wb.create_sheet("Instructions")
    title_font = Font(bold=True, size=14, color="0F172A")
    head_font = Font(bold=True, size=11, color="10B981")
    ins["A1"] = "How to bulk-upload drug categories"
    ins["A1"].font = title_font
    ins.column_dimensions["A"].width = 24
    ins.column_dimensions["B"].width = 72
    row = 3
    ins.cell(row=row, column=1, value="Steps").font = head_font
    for line in (
        "1. Fill one row per category on the 'Categories' sheet.",
        "2. Category Name (*) is required and must be unique.",
        "3. Category Code is optional but, if given, must also be unique.",
        "4. Save the file, then upload it back in the app.",
        "5. Any rows with problems are reported by row number — fix and re-upload only those.",
    ):
        row += 1
        ins.cell(row=row, column=2, value=line)
    row += 2
    ins.cell(row=row, column=1, value="Columns").font = head_font
    for header, _key, required, help_text in CATEGORY_COLUMNS:
        row += 1
        label = ins.cell(row=row, column=1, value=f"{header}{' *' if required else ''}")
        label.font = Font(bold=required)
        ins.cell(row=row, column=2, value=help_text)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _normalize_category_header(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip().lower().rstrip("*").strip()
    return _CATEGORY_HEADER_ALIASES.get(text)


def parse_category_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    """Parse an uploaded category template into ``[{"row", "data"}]`` dicts."""
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb["Categories"] if "Categories" in wb.sheetnames else wb[wb.sheetnames[0]]

    rows_iter = ws.iter_rows(values_only=False)
    try:
        header_cells = next(rows_iter)
    except StopIteration:
        return []

    col_map: dict[int, str] = {}
    for cell in header_cells:
        key = _normalize_category_header(cell.value)
        if key:
            col_map[cell.column] = key
    if "name" not in col_map.values():
        raise ValueError(
            "The sheet doesn't look like the category template (missing the 'Category Name' column)."
        )

    parsed: list[dict[str, Any]] = []
    for cells in rows_iter:
        data: dict[str, Any] = {}
        for cell in cells:
            key = col_map.get(cell.column)
            if not key:
                continue
            cleaned = _clean_cell(cell.value)
            if cleaned is not None:
                data[key] = cleaned
        if not data:
            continue
        parsed.append({"row": cells[0].row if cells else None, "data": data})
    return parsed
